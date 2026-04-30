from __future__ import annotations

import logging
import re
from typing import Any, Optional

from vs_app.integrations.llm.client import complete_text
from vs_app.modules.tickets.documents import ChunkDocument
from vs_app.modules.tickets.text_formatting import (
    clean_text,
    extract_description_text,
    extract_substantive_comments,
)

from ..attachments.materializer import materialize_attachment_documents
from ..attachments.roles import doc_role_weight, normalize_doc_role
from ..attachments.router import get_routing_candidates, route_attachments

logger = logging.getLogger(__name__)

_BODY_PARAGRAPH_MIN_WORDS = 15
_MAX_BODY_PARAGRAPHS = 40
_PREFETCHABLE_EXTENSIONS = {
    "pptx", "ppt", "pdf", "docx", "doc", "xlsx", "xls", "csv",
    "png", "jpg", "jpeg", "gif", "bmp", "tiff", "webp",
}
_IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "bmp", "tiff", "webp"}
_SPREADSHEET_EXTENSIONS = {"xlsx", "xls", "csv"}
_GENERIC_CHUNK_TARGET_WORDS = 425
_GENERIC_CHUNK_MIN_WORDS = 8


async def build_from_attachments(
    *,
    ticket_key: str,
    ticket_data: dict,
    jira_client: Any,
    cfg: Any,
) -> tuple[list[ChunkDocument], list[ChunkDocument], dict[str, Any]]:
    attachments = list(ticket_data.get("attachments", []) or [])
    debug: dict[str, Any] = {
        "attachment_count_total": len(attachments),
        "prefetched_attachment_ids": [],
        "prefetched_attachment_count": 0,
        "used_attachment_ids": [],
        "used_attachment_count": 0,
        "routing": {},
        "attachment_attempts": [],
    }
    if not attachments:
        return [], [], debug

    fields = ticket_data.get("fields", {}) or {}
    ticket_summary = str(fields.get("summary") or ticket_key)

    seed_primary, _seed_supporting, _seed_quality, seed_artifact = route_attachments(
        attachments=attachments,
        ticket_summary=ticket_summary,
        download_fn=None,
    )
    prefetch_targets = unique_attachments(
        ([seed_primary] if seed_primary else [])
        + get_routing_candidates(seed_artifact)
        + attachments
    )
    prefetched = await prefetch_attachment_bytes(jira_client, prefetch_targets, cfg)
    debug["prefetched_attachment_ids"] = list(prefetched.keys())
    debug["prefetched_attachment_count"] = len(prefetched)

    if prefetched:
        def cached_download(att: dict) -> bytes:
            att_id = attachment_key(att)
            file_bytes = prefetched.get(att_id)
            if file_bytes is None:
                raise RuntimeError(f"Not prefetched: {att.get('filename')}")
            return file_bytes

        primary, _supporting, _quality, routing_artifact = route_attachments(
            attachments=attachments,
            ticket_summary=ticket_summary,
            download_fn=cached_download,
        )
    else:
        primary = seed_primary
        routing_artifact = seed_artifact

    debug["routing"] = _summarize_routing_artifact(routing_artifact)
    chunk_candidates = get_routing_candidates(routing_artifact)
    ordered_candidates = unique_attachments(([primary] if primary else []) + chunk_candidates + attachments)
    planned_chunk_ids = {
        str(value).strip()
        for value in (
            (routing_artifact.get("processing_plan") or {}).get("chunk_candidates")
            or []
        )
        if str(value).strip()
    }

    processed_docs: list[dict[str, Any]] = []
    max_docs = _optional_int_limit(getattr(cfg, "max_chunk_attachments", 6))
    for att in ordered_candidates:
        if max_docs is not None and len(debug["attachment_attempts"]) >= max_docs:
            break

        att_id = attachment_key(att)
        att_debug: dict[str, Any] = {
            "attachment_id": att_id,
            "filename": str(att.get("filename") or att_id),
            "attachment_type": str(att.get("ext") or ext_of(att.get("filename", ""))),
            "doc_role": normalize_doc_role(att),
            "triage_score": att.get("triage_score"),
            "selected_for_chunking": att_id in planned_chunk_ids,
            "download_status": "not_started",
            "extraction_status": "not_started",
            "leaf_count": 0,
            "word_count": 0,
            "error": "",
            "used_for_chunking": False,
        }

        file_bytes = prefetched.get(att_id)
        if file_bytes is None:
            try:
                file_bytes = await jira_client.download_attachment(att)
                att_debug["download_status"] = "downloaded_on_demand"
            except Exception as exc:
                logger.info("Attachment download skipped (%s): %s", att.get("filename"), exc)
                att_debug["download_status"] = "download_failed"
                att_debug["error"] = str(exc)
                debug["attachment_attempts"].append(att_debug)
                continue
        else:
            att_debug["download_status"] = "prefetched"

        extraction = extract_attachment_payload(file_bytes or b"", att, cfg)
        leaves = [leaf for leaf in (extraction.get("leaves") or []) if leaf.get("text") and not leaf.get("is_boilerplate")]
        att_debug["attachment_type"] = str(extraction.get("ext") or att_debug["attachment_type"])
        att_debug["extraction_status"] = str(extraction.get("status") or "unknown")
        att_debug["word_count"] = int(extraction.get("word_count") or 0)
        att_debug["leaf_count"] = len(leaves)
        att_debug["error"] = str(extraction.get("error") or "")
        if _attachment_debug_enabled(cfg):
            att_debug["raw_extracted_text"] = str(extraction.get("raw_extracted_text") or "")
            att_debug["leaf_chunks"] = leaves
        if not leaves:
            debug["attachment_attempts"].append(att_debug)
            continue

        processed_docs.append(
            {
                "attachment_id": att_id,
                "attachment_name": str(att.get("filename") or att_id),
                "attachment_type": str(extraction.get("ext") or att.get("ext") or ext_of(att.get("filename", ""))),
                "source_url": str(att.get("content") or att.get("url") or ticket_key),
                "doc_role": normalize_doc_role(att),
                "leaves": leaves,
            }
        )
        att_debug["used_for_chunking"] = True
        debug["used_attachment_ids"].append(att_id)
        debug["attachment_attempts"].append(att_debug)

    if not processed_docs:
        return [], [], debug

    debug["used_attachment_count"] = len(debug["used_attachment_ids"])
    sections, leaves = materialize_attachment_documents(ticket_key, processed_docs)
    return sections, leaves, debug


async def prefetch_attachment_bytes(
    jira_client: Any,
    attachments: list[dict],
    cfg: Any,
) -> dict[str, bytes]:
    max_prefetch = _optional_int_limit(getattr(cfg, "max_prefetch_attachments", 8))
    max_size = int(getattr(cfg, "max_prefetch_attachment_size", 60_000_000) or 60_000_000)

    prefetched: dict[str, bytes] = {}
    for att in attachments:
        if max_prefetch is not None and len(prefetched) >= max_prefetch:
            break
        ext = ext_of(att.get("filename", ""))
        size = int(att.get("size", 0) or 0)
        if ext not in _PREFETCHABLE_EXTENSIONS:
            continue
        if size and size > max_size:
            continue

        att_id = attachment_key(att)
        if att_id in prefetched:
            continue
        try:
            prefetched[att_id] = await jira_client.download_attachment(att)
        except Exception as exc:
            logger.info("Attachment prefetch skipped (%s): %s", att.get("filename"), exc)

    return prefetched


def extract_attachment_payload(file_bytes: bytes, attachment: dict, cfg: Any) -> dict[str, Any]:
    ext = str(attachment.get("ext") or ext_of(attachment.get("filename", ""))).lower()
    filename = str(attachment.get("filename") or f"attachment.{ext or 'bin'}")
    extracted = attachment.get("extracted") or {}
    raw_extracted_text = ""

    try:
        if not extracted.get("chunks"):
            if ext in ("pptx", "ppt"):
                from vs_app.integrations.files.pptx_extractor import extract_pptx

                extracted = extract_pptx(
                    file_bytes,
                    max_slides=int(getattr(cfg, "max_slides", 60) or 60),
                )
            elif ext == "pdf":
                from vs_app.integrations.files.pdf_extractor import extract_pdf

                extracted = extract_pdf(
                    file_bytes,
                    ocr_enabled=bool(getattr(cfg, "ocr_enabled", False)),
                    max_pages=int(getattr(cfg, "max_pages", 60) or 60),
                )
            elif ext in ("docx", "doc"):
                from vs_app.integrations.files.docx_extractor import extract_docx

                extracted = extract_docx(file_bytes)
    except Exception as exc:
        logger.warning("Attachment extraction failed (%s): %s", attachment.get("filename"), exc)
        return {
            "ext": ext,
            "status": "extract_failed",
            "error": str(exc),
            "word_count": 0,
            "raw_extracted_text": "",
            "leaves": [],
        }

    raw_extracted_text = _combine_extracted_chunks(extracted)
    if (not extracted.get("chunks")) or _attachment_debug_enabled(cfg):
        if not raw_extracted_text:
            raw_extracted_text = _extract_raw_attachment_text(
                file_bytes=file_bytes,
                filename=filename,
                ext=ext,
                ocr_enabled=bool(getattr(cfg, "ocr_enabled", False)),
            )
    if not extracted.get("chunks") and raw_extracted_text:
        extracted = {
            "chunks": _build_generic_text_chunks(raw_extracted_text, attachment, ext),
        }

    leaves = _convert_extracted_chunks_to_leaves(extracted, attachment, ext)
    status = "extracted" if leaves else "no_text"
    return {
        "ext": ext,
        "status": status,
        "error": "",
        "word_count": len((raw_extracted_text or "").split()) if raw_extracted_text else sum(
            int(leaf.get("word_count") or 0) for leaf in leaves
        ),
        "raw_extracted_text": raw_extracted_text,
        "leaves": leaves,
    }


def extract_attachment_leaves(file_bytes: bytes, attachment: dict, cfg: Any) -> list[dict]:
    return list(extract_attachment_payload(file_bytes, attachment, cfg).get("leaves") or [])


def _convert_extracted_chunks_to_leaves(
    extracted: dict[str, Any],
    attachment: dict,
    ext: str,
) -> list[dict]:
    doc_role = normalize_doc_role(attachment)
    weight = doc_role_weight(doc_role)

    leaves: list[dict] = []
    for idx, raw in enumerate(extracted.get("chunks", []) or []):
        text = clean_text(str(raw.get("text") or ""))
        if not text:
            continue
        section_title = str(
            raw.get("section_title")
            or raw.get("slide_title")
            or attachment.get("filename")
            or ""
        ).strip()
        slide_num = raw.get("slide_num")
        page_num = raw.get("page_num")
        page_range = raw.get("page_range")
        slide_range = raw.get("slide_range")
        source = str(raw.get("source") or source_for_ext(ext))
        header_hierarchy = build_header_hierarchy(
            section_title,
            slide_num,
            page_num,
            page_range,
            slide_range,
        )

        leaves.append(
            {
                "chunk_id": str(raw.get("chunk_id") or f"{attachment_key(attachment)}-{idx}"),
                "text": text,
                "source": source,
                "section_title": section_title,
                "slide_title": str(raw.get("slide_title") or ""),
                "slide_num": slide_num,
                "page_num": page_num,
                "page_range": page_range,
                "slide_range": slide_range,
                "is_boilerplate": bool(raw.get("is_boilerplate", False)),
                "weight_multiplier": float(raw.get("weight_multiplier", weight) or weight),
                "extraction_confidence": float(raw.get("extraction_confidence", 0.9) or 0.9),
                "word_count": int(raw.get("word_count") or len(text.split())),
                "attachment_id": attachment_key(attachment),
                "attachment_name": str(attachment.get("filename") or ""),
                "attachment_type": ext,
                "header_hierarchy": header_hierarchy,
            }
        )

    return leaves


def build_from_ticket_body(
    ticket_key: str,
    ticket_data: dict,
    llm_client: Optional[Any],
    cfg: Any,
) -> tuple[list[ChunkDocument], list[ChunkDocument]]:
    fields = ticket_data.get("fields", {}) or {}
    desc_text = clean_text(extract_description_text(fields.get("description")))
    comments = extract_substantive_comments(fields.get("comment") or {}, max_comments=3)

    if not desc_text and not comments:
        return [], []

    body_parts: list[tuple[str, str]] = []
    for paragraph in split_paragraphs(desc_text):
        body_parts.append(("description", paragraph))
    for comment in comments:
        for paragraph in split_paragraphs(comment):
            body_parts.append(("comment", paragraph))
    body_parts = body_parts[:_MAX_BODY_PARAGRAPHS]

    if not body_parts:
        return [], []

    overview = generate_body_overview(ticket_key, desc_text, comments, llm_client, cfg)
    title = str(fields.get("summary") or ticket_key)
    section_text = f"## Ticket Body: {title}\n\n{overview}" if overview else f"## Ticket Body: {title}"

    attachment_id = ""
    section_uid = ChunkDocument.make_uid(ticket_key, attachment_id, "section", "ticket-body")

    section_doc = ChunkDocument(
        chunk_uid=section_uid,
        ticket_id=ticket_key,
        level="section",
        parent_uid="",
        attachment_id=attachment_id,
        section_uid=section_uid,
        text=section_text,
        source="section",
        section_title=title,
        doc_role="ticket_body",
        weight_multiplier=0.9,
        extraction_confidence=1.0,
        word_count=len(section_text.split()),
    )

    leaf_docs: list[ChunkDocument] = []
    for idx, (source, text) in enumerate(body_parts):
        leaf_docs.append(
            ChunkDocument(
                chunk_uid=ChunkDocument.make_uid(ticket_key, attachment_id, source, f"body-{idx}"),
                ticket_id=ticket_key,
                level="chunk",
                parent_uid=section_uid,
                attachment_id=attachment_id,
                section_uid=section_uid,
                text=text,
                source=source,
                section_title=title,
                doc_role="description" if source == "description" else "comment",
                chunk_index=idx,
                weight_multiplier=0.9 if source == "description" else 0.6,
                extraction_confidence=1.0,
                word_count=len(text.split()),
            )
        )

    return [section_doc], leaf_docs


def split_paragraphs(text: str) -> list[str]:
    if not text:
        return []
    paragraphs = re.split(r"\n{2,}", text)
    return [paragraph.strip() for paragraph in paragraphs if len(paragraph.split()) >= _BODY_PARAGRAPH_MIN_WORDS]


def generate_body_overview(
    ticket_key: str,
    description: str,
    comments: list[str],
    llm_client: Optional[Any],
    cfg: Any,
) -> str:
    if llm_client is None or getattr(cfg, "skip_llm_summary", False):
        return ""
    if not description and not comments:
        return ""

    body = description
    if comments:
        body = body + "\n\n" + "\n\n".join(comments)
    body = body[:6000]

    prompt = (
        "Write a 2-3 sentence plain-language overview of the following Jira ticket "
        "body. Focus on what business problem is being described and what capability "
        "or change is being requested.\n\n"
        f"Ticket {ticket_key} body:\n{body}"
    )
    try:
        return clean_text(
            complete_text(
                prompt,
                llm_client,
                model=getattr(cfg, "llm_model", None),
                max_output_tokens=400,
                temperature=0.2,
            )
            or ""
        )
    except Exception as exc:
        logger.info("Body overview generation failed: %s", exc)
        return ""


def unique_attachments(attachments: list[Optional[dict]]) -> list[dict]:
    out: list[dict] = []
    seen: set[str] = set()
    for att in attachments:
        if not att:
            continue
        att_id = attachment_key(att)
        if not att_id or att_id in seen:
            continue
        seen.add(att_id)
        out.append(att)
    return out


def attachment_key(att: dict) -> str:
    return str(att.get("id") or att.get("attachment_id") or att.get("filename") or "")


def ext_of(filename: str) -> str:
    name = (filename or "").lower()
    return name.rsplit(".", 1)[-1] if "." in name else ""


def source_for_ext(ext: str) -> str:
    if ext in {"pptx", "ppt"}:
        return "pptx_slide"
    if ext == "pdf":
        return "pdf_page"
    return "docx_section"


def build_header_hierarchy(
    section_title: str,
    slide_num: Optional[int],
    page_num: Optional[int],
    page_range: Optional[list[int]],
    slide_range: Optional[list[int]],
) -> str:
    parts: list[str] = []
    if section_title:
        parts.append(section_title)
    locator = locator_label(
        slide_num=slide_num,
        page_num=page_num,
        page_range=page_range,
        slide_range=slide_range,
    )
    if locator:
        parts.append(locator)
    return " > ".join(parts)


def locator_label(
    *,
    slide_num: Optional[int] = None,
    page_num: Optional[int] = None,
    page_range: Optional[list[int]] = None,
    slide_range: Optional[list[int]] = None,
) -> str:
    if page_num is not None:
        return f"page:{page_num}"
    if slide_num is not None:
        return f"slide:{slide_num}"
    if page_range:
        return f"pages:{page_range[0]}-{page_range[-1]}"
    if slide_range:
        return f"slides:{slide_range[0]}-{slide_range[-1]}"
    return ""


def _attachment_debug_enabled(cfg: Any) -> bool:
    return bool(
        getattr(cfg, "enable_raw_artifact_persistence", False)
        or getattr(cfg, "enable_attachment_text_persistence", False)
        or getattr(cfg, "enable_debug_stage_persistence", False)
        or getattr(cfg, "enable_prechunk_persistence", False)
    )


def _optional_int_limit(value: Any) -> Optional[int]:
    if value in (None, "", 0, "0", False):
        return None
    try:
        return max(int(value), 1)
    except (TypeError, ValueError):
        return None


def _combine_extracted_chunks(extracted: dict[str, Any]) -> str:
    parts = [str(chunk.get("text") or "").strip() for chunk in (extracted.get("chunks") or [])]
    return "\n\n".join(part for part in parts if part).strip()


def _extract_raw_attachment_text(
    *,
    file_bytes: bytes,
    filename: str,
    ext: str,
    ocr_enabled: bool,
) -> str:
    try:
        if ext in _IMAGE_EXTENSIONS:
            from vs_app.integrations.files.markitdown_extractor import extract_image_text

            return clean_text(extract_image_text(file_bytes, filename=filename))

        from vs_app.integrations.files.markitdown_extractor import extract_markdown

        raw = extract_markdown(
            file_bytes,
            filename,
            enable_ocr=ocr_enabled,
        )
        if raw:
            return clean_text(raw)
    except Exception as exc:
        logger.info("Raw attachment extraction fallback failed (%s): %s", filename, exc)

    if ext == "csv":
        return clean_text(file_bytes.decode("utf-8", errors="replace"))
    if ext in _SPREADSHEET_EXTENSIONS:
        return clean_text(_extract_spreadsheet_text(file_bytes))
    return ""


def _extract_spreadsheet_text(file_bytes: bytes) -> str:
    try:
        import io
        import openpyxl
    except Exception:
        return ""

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        return ""

    parts: list[str] = []
    for worksheet in workbook.worksheets:
        parts.append(f"## {worksheet.title}")
        for row in worksheet.iter_rows(values_only=True):
            cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if cells:
                parts.append(" | ".join(cells))
        parts.append("")
    return "\n".join(parts).strip()


def _build_generic_text_chunks(raw_text: str, attachment: dict, ext: str) -> list[dict]:
    cleaned = clean_text(raw_text)
    if not cleaned:
        return []

    sections = _split_generic_sections(cleaned, fallback_title=str(attachment.get("filename") or "Attachment"))
    chunks: list[dict] = []
    chunk_idx = 0
    for section in sections:
        paragraphs = [part.strip() for part in re.split(r"\n{2,}", section["body"]) if part.strip()]
        current: list[str] = []
        current_words = 0

        def flush() -> None:
            nonlocal chunk_idx, current, current_words
            if not current:
                return
            text = "\n\n".join(current).strip()
            words = len(text.split())
            if words >= _GENERIC_CHUNK_MIN_WORDS:
                chunks.append(
                    {
                        "chunk_id": f"{attachment_key(attachment)}-generic-{chunk_idx}",
                        "source": source_for_ext(ext),
                        "text": text,
                        "section_title": section["title"],
                        "word_count": words,
                        "is_boilerplate": False,
                        "weight_multiplier": 1.0,
                        "extraction_confidence": 0.75,
                        "extraction_method": "markitdown_generic",
                    }
                )
                chunk_idx += 1
            current = []
            current_words = 0

        for paragraph in paragraphs:
            words = len(paragraph.split())
            if current and current_words + words > _GENERIC_CHUNK_TARGET_WORDS:
                flush()
            current.append(paragraph)
            current_words += words
        flush()

    if not chunks and cleaned.strip():
        chunks.append(
            {
                "chunk_id": f"{attachment_key(attachment)}-generic-0",
                "source": source_for_ext(ext),
                "text": cleaned,
                "section_title": str(attachment.get("filename") or "Attachment"),
                "word_count": len(cleaned.split()),
                "is_boilerplate": False,
                "weight_multiplier": 1.0,
                "extraction_confidence": 0.7,
                "extraction_method": "markitdown_generic",
            }
        )
    return chunks


def _split_generic_sections(text: str, fallback_title: str) -> list[dict[str, str]]:
    matches = list(re.finditer(r"(?m)^(#{1,6})\s+(.+)$", text))
    if not matches:
        return [{"title": fallback_title, "body": text}]

    sections: list[dict[str, str]] = []
    preamble = text[:matches[0].start()].strip()
    if preamble:
        sections.append({"title": fallback_title, "body": preamble})

    for idx, match in enumerate(matches):
        title = match.group(2).strip() or fallback_title
        start = match.end()
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
        body = text[start:end].strip()
        if body:
            sections.append({"title": title, "body": body})
    return sections


def _summarize_routing_artifact(artifact: dict[str, Any]) -> dict[str, Any]:
    processing_plan = dict(artifact.get("processing_plan") or {})
    return {
        "att_quality": artifact.get("att_quality"),
        "quality_tier": artifact.get("quality_tier"),
        "selection_reason": artifact.get("selection_reason"),
        "attachment_count_total": artifact.get("attachment_count_total"),
        "attachment_count_viable": artifact.get("attachment_count_viable"),
        "primary_attachment_id": artifact.get("primary_attachment_id"),
        "triage_score": artifact.get("triage_score"),
        "triage_reasons": list(artifact.get("triage_reasons") or []),
        "primary_attachment": artifact.get("primary_attachment"),
        "supporting_attachments": list(artifact.get("supporting_attachments") or []),
        "excluded_attachments": list(artifact.get("excluded_attachments") or []),
        "per_attachment_scores": list(artifact.get("per_attachment_scores") or []),
        "processing_plan": {
            "full_extract": list(processing_plan.get("full_extract") or []),
            "light_extract": list(processing_plan.get("light_extract") or []),
            "skip": list(processing_plan.get("skip") or []),
            "chunk_candidates": list(processing_plan.get("chunk_candidates") or []),
        },
    }


__all__ = [
    "attachment_key",
    "build_from_attachments",
    "build_from_ticket_body",
    "extract_attachment_leaves",
    "extract_attachment_payload",
    "ext_of",
    "prefetch_attachment_bytes",
    "unique_attachments",
]
